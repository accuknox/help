"""Weekly competitor-tool signal: job post that names a rival tool -> company -> buyers -> verified emails.

Firecrawl finds and reads the job pages. treg enriches companies, finds people and emails.
Jev judges every fuzzy call. See ../SKILL.md. Run:

    python run.py --target 25 --cap 0.60
"""
import argparse
import concurrent.futures as cf
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import urlparse

sys.path.insert(0, r"D:\Atharva\NOTES\SCRIPTS\gtm")
from common import Ledger, output, stamp  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

HERE = Path(__file__).resolve().parent
STATE = HERE.parent / ".state"
OUT = Path(r"D:\AccuKnox\help\drafts\competitor-signal")
STATE.mkdir(exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

# tool -> (Firecrawl query, regex that must appear in the page)
TOOLS = {
    "Wiz": ('"Wiz" cloud security engineer job', r"\bWiz\b"),
    "Orca": ('"Orca Security" cloud security job', r"\bOrca\b"),
    "Lacework": ('"Lacework" security engineer job', r"\bLacework\b"),
    "Aqua": ('"Aqua Security" container security job', r"\bAqua\b"),
    "Sysdig": ('"Sysdig" security engineer job', r"\bSysdig\b"),
    "Prisma Cloud": ('"Prisma Cloud" security engineer job', r"\bPrisma\b"),
    "Veracode": ('"Veracode" application security engineer job', r"\bVeracode\b"),
    "Checkmarx": ('"Checkmarx" application security engineer job', r"\bCheckmarx\b"),
}
# Extra queries for a deep hunt (--variants). Each reads as a different way a job post names the tool.
VARIANTS = {
    "Prisma Cloud": ['"Prisma Cloud" CSPM Palo Alto security engineer', '"Prisma Cloud Compute" OR Twistlock container security', 'Palo Alto "Prisma Cloud" DevSecOps careers'],
    "Orca": ['"Orca Security" CNAPP cloud engineer', 'Orca agentless cloud security posture engineer'],
    "Sysdig": ['"Sysdig Secure" Falco Kubernetes runtime security', 'Sysdig container runtime security engineer careers'],
}
OPENER = ("Saw {company} is hiring for a role that works with {tool}. AccuKnox adds runtime protection for "
          "cloud, Kubernetes and AI workloads on top of it. Open to a short chat?")

PRODUCT = ("AccuKnox is a cloud and AI security platform: CNAPP (cloud posture, Kubernetes and container "
           "runtime protection, CI/CD and IaC scanning), AI and LLM security, and zero trust. It is sold to "
           "security leaders at mid-size and large enterprises.")
TIERS = {
    "high":      {"buyer": 0.75, "offer": 0.60, "competitor_max": 0.25, "can_buy": 0.70, "evidence": 0.80},
    "qualified": {"buyer": 0.55, "offer": 0.50, "competitor_max": 0.40, "can_buy": 0.50, "evidence": 0.60},
}
FIT = [{"not a lead": "recruiters, HR, sales, students, finance, or anyone whose work is not security or cloud infrastructure"},
       {"practitioner": "a security or DevOps engineer or analyst who feels the problem but does not select tools"},
       {"decision maker": "a CISO, VP, head or director of security, cloud security, product security or DevSecOps, "
                          "or a principal or staff security architect who selects or approves security tools"}]
BUYS = {"type": "choice", "instructions": "Which offer fits this person best? `product` is AccuKnox, the seller.",
        "criteria": {"accuknox": "a security leader or architect who would evaluate a cloud, container or AI security platform",
                     "neither": "not a buyer of cloud or AI security tooling"}}
ROLE = {"type": "choice", "instructions": "What is this person's role?",
        "criteria": {"ciso_exec": "CISO, VP or head of security or infrastructure",
                     "sec_director": "director or senior manager of security, cloud security, product security or DevSecOps",
                     "architect": "security architect or principal or staff engineer",
                     "engineer": "security, cloud or DevOps engineer or analyst", "other": "anything else"}}
PEOPLE_QS = {
    "fit": {"type": "score", "instructions": "How good a sales lead is this person for `product`? Judge by the role in `position`.", "criteria": FIT},
    "role": ROLE, "buys": BUYS,
    "competitor": {"type": "noul", "instructions": "Does this person work at a company that sells cloud, container or AI security software? `product` is the seller, never this person's.",
                   "criteria": {"true": "their title or company sells cybersecurity or cloud security products or services",
                                "false": "they work in-house at a company that would buy security tooling"}},
    "can_buy": {"type": "noul", "instructions": "Can this person select or approve a security platform purchase, given the role and the company size?"},
    "evidence": {"type": "noul", "instructions": "Does the `position` text itself name a security leadership role such as CISO, VP, head, director, architect or principal? A vague or empty headline is no."}}


def to_int(x):
    """Headcount from a number or a band such as '10K+', 'over-10K' or '1001-5000'."""
    m = re.findall(r"(\d[\d,]*\.?\d*)\s*([kK]?)", str(x))
    if not m:
        return 0
    n, k = m[0]
    return int(float(n.replace(",", "")) * (1000 if k else 1))


def same_org(email, domain):
    e, c = email.split("@")[1].lower().split("."), domain.lower().split(".")
    return len(e) > 1 and len(c) > 1 and e[-2][:5] == c[-2][:5]


def qualify(w):
    if w.get("p_buyer") is None:
        return None, "jev failed"
    for tier, t in TIERS.items():
        if (w["p_buyer"] >= t["buyer"] and w["p_offer"] >= t["offer"] and w["p_competitor"] <= t["competitor_max"]
                and w["p_can_buy"] >= t["can_buy"] and w["p_evidence"] >= t["evidence"]):
            return tier, ""
    return None, "failed a check"


def company_hint(url, title):
    """Best guess at the employer name from the title or the URL. Jev checks it later."""
    m = re.search(r"\bat ([A-Z][^|\u2014\u2013\-–]{1,40}?)(?:\s[-|\u2014\u2013]|\s\(|$)", title or "")
    if m:
        return m.group(1).strip()
    u = urlparse(url)
    parts = [p for p in u.path.split("/") if p]
    if any(h in u.netloc for h in ("greenhouse.io", "lever.co", "ashbyhq.com", "smartrecruiters.com", "workable.com")):
        return parts[0] if parts else ""
    if "myworkdayjobs.com" in u.netloc:
        return u.netloc.split(".")[0]
    if u.netloc.startswith(("jobs.", "careers.")):
        return u.netloc.split(".")[1]
    return ""


def tbs_for(window):
    """w or m for the last week or month, or a number of days such as 60 for a custom range."""
    if window.isdigit():
        import datetime as dt
        end = dt.date.today()
        start = end - dt.timedelta(days=int(window))
        return f"cdr:1,cd_min:{start.month}/{start.day}/{start.year},cd_max:{end.month}/{end.day}/{end.year}"
    return f"qdr:{window}"


def firecrawl_search(query, limit, window="w"):
    r = subprocess.run([shutil.which("firecrawl") or "firecrawl", "search", query, "--tbs", tbs_for(window), "--limit", str(limit), "--scrape", "--json"],
                       capture_output=True, text=True, encoding="utf-8", timeout=240)
    i = r.stdout.find("{")
    try:
        d = json.loads(r.stdout[i:])
    except ValueError:
        return []
    data = d.get("data", d)
    return data.get("web", []) if isinstance(data, dict) else []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--target", type=int, default=25)
    ap.add_argument("--cap", type=float, default=0.60)
    ap.add_argument("--per-tool", type=int, default=8)
    ap.add_argument("--window", default="w", help="w = last week, m = last month, or a number of days such as 60")
    ap.add_argument("--tools", default="", help="comma list of tools to watch, default all")
    ap.add_argument("--variants", action="store_true", help="also run the extra query variants for each tool")
    ap.add_argument("--no-seen", action="store_true", help="ignore the seen list")
    ap.add_argument("--min-employees", type=int, default=1000)
    ap.add_argument("--per-company", type=int, default=2)
    a = ap.parse_args()
    L = Ledger(a.cap)
    seen_f = STATE / "seen.json"
    seen = set(json.loads(seen_f.read_text())) if seen_f.exists() and not a.no_seen else set()
    run = {}

    # 1. Firecrawl finds the job pages from the last week and reads them.
    pages = []
    want = [t.strip() for t in a.tools.split(",") if t.strip()]
    jobs = [(t, q, rx) for t, (q, rx) in TOOLS.items() if not want or t in want]
    if a.variants:
        jobs += [(t, q, TOOLS[t][1]) for t, qs in VARIANTS.items() if not want or t in want for q in qs]
    seen_urls = set()
    for tool, q, rx in jobs:
        for r in firecrawl_search(q, a.per_tool, a.window):
            if r.get("url") in seen_urls:
                continue
            seen_urls.add(r.get("url"))
            md = r.get("markdown") or ""
            m = re.search(rx, md)
            if not m or "linkedin.com" in r.get("url", ""):
                continue
            title = (r.get("metadata") or {}).get("title") or r.get("title") or ""
            pages.append({"tool": tool, "url": r["url"], "title": title, "hint": company_hint(r["url"], title),
                          "snippet": md[max(0, m.start() - 140): m.end() + 140].replace("\n", " ")})
    run["pages_with_tool_name"] = len(pages)
    print(f"1. {len(pages)} job pages name a tool")

    # 2. Jev: is it a real job post where the team uses the tool, and is the employer a plain company?
    def judge_page(p):
        r = L.jev({"tool": p["tool"], "snippet": p["snippet"], "page_title": p["title"], "url": p["url"]}, {
            "uses": {"type": "noul", "instructions": "Does `snippet` say the team uses, runs, administers or wants experience with the security product named in `tool`? The word with another meaning, the vendor itself hiring, or a staffing ad is no."},
            "job": {"type": "noul", "instructions": "Is this page a single job posting for a security, cloud or DevOps role?"}})
        p["p_uses"], p["p_job"] = (r["answers"]["uses"]["noul"], r["answers"]["job"]["noul"]) if r else (None, None)
        return p
    with cf.ThreadPoolExecutor(6) as ex:
        pages = list(ex.map(judge_page, pages))
    good = [p for p in pages if p["p_uses"] is not None and p["p_uses"] >= 0.6 and p["p_job"] >= 0.6 and p["hint"]]
    run["jobs_confirmed"] = len(good)
    print(f"2. {len(good)} confirmed job posts")

    # 3. treg names the company, Jev confirms the match, size gate.
    companies = {}
    for p in good:
        key = p["hint"].lower()
        if key in companies:
            companies[key]["tools"].add(p["tool"]); continue
        o = None
        for guess in (p["hint"], re.split(r"\s[|—–-]\s", p["title"])[-1].strip()):
            d = L.treg_call("treg.companies.enrich", body={"name": guess}, max_cost=0.01)
            o = output(d) if d else None
            if o and o.get("domain"):
                break
        if not (o and o.get("domain")):
            continue
        r = L.jev({"guess": p["hint"], "enriched_name": o.get("name"), "domain": o["domain"], "page_title": p["title"]},
                  {"same": {"type": "noul", "instructions": "Could `guess` refer to `enriched_name`? A short name, an abbreviation or a name without spaces counts as the same company."},
                   "vendor": {"type": "noul", "instructions": "Is `enriched_name` a cybersecurity software vendor, or a staffing and recruiting agency? Banks, IT services firms, consultancies and government contractors are no."}})
        if not r or r["answers"]["same"]["noul"] < 0.6 or r["answers"]["vendor"]["noul"] >= 0.5:
            continue
        companies[key] = {"name": o["name"], "domain": o["domain"], "employees": to_int(o.get("employees")),
                          "industry": o.get("industry"), "tools": {p["tool"]}, "job": p["title"], "job_url": p["url"]}
    sized = {k: c for k, c in companies.items() if c["employees"] >= a.min_employees}
    new = {k: c for k, c in sized.items() if not all(f'{c["domain"]}|{t}' in seen for t in c["tools"])}
    run.update(companies=len(companies), companies_sized=len(sized), companies_new=len(new))
    print(f"3. {len(companies)} companies, {len(sized)} sized, {len(new)} new this week")

    # 4. treg finds security people, Jev runs the five checks.
    people = {}
    for c in new.values():
        d = L.treg_call("treg.people.search", body={"company_domain": c["domain"], "title": "security", "limit": 25}, max_cost=0.05)
        for p in (output(d) or {}).get("people", []) if d else []:
            if p.get("first_name") and p.get("title") and p.get("employee_linkedin"):
                people[p["employee_linkedin"]] = {
                    "name": f'{p["first_name"]} {p["last_name"]}'.strip(), "first": p["first_name"], "last": p["last_name"],
                    "title": p["title"], "linkedin": p["employee_linkedin"], "company": c["name"], "domain": c["domain"],
                    "employees": c["employees"], "industry": c["industry"], "tool": ", ".join(sorted(c["tools"])),
                    "job": c["job"], "job_url": c["job_url"]}

    def judge(item):
        pid, w = item
        r = L.jev({"position": w["title"], "company": w["company"], "industry": w["industry"], "employees": w["employees"], "product": PRODUCT}, PEOPLE_QS)
        if r:
            A = r["answers"]
            w.update(p_buyer=A["fit"]["probabilities"].get("2", 0), p_offer=A["buys"]["probabilities"].get("accuknox", 0),
                     p_competitor=A["competitor"]["noul"], p_can_buy=A["can_buy"]["noul"], p_evidence=A["evidence"]["noul"])
        else:
            w["p_buyer"] = None
        w["tier"], _ = qualify(w)
        return pid, w
    with cf.ThreadPoolExecutor(6) as ex:
        people = dict(ex.map(judge, people.items()))
    passed = sorted((w for w in people.values() if w["tier"]),
                    key=lambda w: (w["tier"] != "high", -(w["p_buyer"] + w["p_offer"] + w["p_can_buy"] + w["p_evidence"])))
    per, chosen = {}, []
    for w in passed:
        if per.get(w["company"], 0) < a.per_company:
            per[w["company"]] = per.get(w["company"], 0) + 1
            chosen.append(w)
    run.update(people=len(people), people_pass=len(passed))
    print(f"4. {len(people)} people, {len(passed)} pass the five checks")

    # 5. treg email waterfall, then verify. Only a valid email on the company's own domain counts.
    leads = []
    for w in chosen:
        if len(leads) >= a.target or L.left() < 0.03:
            break
        d = L.treg_call("treg.people.email.find", body={"first_name": w["first"], "last_name": w["last"], "domain": w["domain"]}, max_cost=0.03)
        em = (output(d) or {}).get("email") if d else None
        if not em:
            continue
        v = L.treg_call("treg.people.email.verify", body={"email": em}, max_cost=0.003)
        if v and (output(v) or {}).get("status") in ("valid", "safe", "deliverable") and same_org(em, w["domain"]):
            w["email"] = em
            leads.append(w)
    run.update(leads=len(leads), treg_usd=round(L.treg, 4), jev_usd=round(L.jev_usd, 5))
    print(f"5. {len(leads)} verified leads · treg ${L.treg:.3f} · jev ${L.jev_usd:.4f}")

    # 6. One sheet, one tab. Remember what was reported so next week shows only new accounts.
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Name", "Title", "Company", "Work email", "Tool they run", "Why now (job post)", "Job link", "LinkedIn", "Opening line"])
    for c in ws[1]:
        c.fill, c.font = PatternFill("solid", fgColor="1D4ED8"), Font(bold=True, color="FFFFFF")
    for w in leads:
        ws.append([w["name"], w["title"], w["company"], w["email"], w["tool"], w["job"], w["job_url"], w["linkedin"],
                   OPENER.format(company=w["company"].rstrip(" ,."), tool=w["tool"].split(", ")[0])])
    for col, wd in zip("ABCDEFGHI", [22, 44, 24, 34, 16, 46, 36, 34, 70]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path = OUT / f"AccuKnox_Competitor_Tool_Leads_{stamp()[:10]}.xlsx"
    wb.save(path)
    for c in new.values():
        seen.update(f'{c["domain"]}|{t}' for t in c["tools"])
    seen_f.write_text(json.dumps(sorted(seen)))
    (STATE / f"run-{stamp()}.json").write_text(json.dumps({"run": run, "leads": leads, "jobs": good}, indent=1, default=str), encoding="utf-8")
    print("saved", path)


if __name__ == "__main__":
    main()
