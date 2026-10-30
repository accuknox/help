"""Merge the old sheets and the hunt output into one deduped workbook.

    python merge.py --out "<path>.xlsx" --old a.xlsx b.xlsx --hunt hunt-*.json
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from run import OPENER  # noqa: E402

COLS = ["Name", "Title", "Company", "Work email", "Tool they run", "Why now (job post)", "Posted", "Job link", "LinkedIn", "Opening line", "Source"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--old", nargs="*", default=[])
    ap.add_argument("--hunt", nargs="*", default=[])
    a = ap.parse_args()
    rows, seen = [], set()

    def add(r):
        k = (r["Work email"] or r["LinkedIn"] or "").lower()
        if k and k not in seen:
            seen.add(k)
            rows.append(r)

    for f in a.old:
        ws = load_workbook(f).active
        head = [c.value for c in ws[1]]
        for v in list(ws.iter_rows(values_only=True))[1:]:
            d = dict(zip(head, v))
            first = d.get("Tool they run") or ""
            add({"Name": d.get("Name"), "Title": d.get("Title"), "Company": d.get("Company"), "Work email": d.get("Work email"),
                 "Tool they run": first, "Why now (job post)": d.get("Why now (job post)"), "Posted": d.get("Posted") or "",
                 "Job link": d.get("Job link"), "LinkedIn": d.get("LinkedIn"),
                 "Opening line": d.get("Opening line") or "", "Source": Path(f).stem})
    for f in a.hunt:
        for w in json.loads(Path(f).read_text(encoding="utf-8"))["leads"]:
            add({"Name": w["name"], "Title": w["title"], "Company": w["company"], "Work email": w["email"], "Tool they run": w["tool"],
                 "Why now (job post)": w["job"], "Posted": w["posted"], "Job link": w["job_url"], "LinkedIn": w["linkedin"],
                 "Opening line": OPENER.format(company=w["company"].rstrip(" ,."), tool=w["tool"].split(", ")[0]), "Source": "TheirStack hunt 2026-10-30"})
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLS)
    for c in ws[1]:
        c.fill, c.font = PatternFill("solid", fgColor="1D4ED8"), Font(bold=True, color="FFFFFF")
    for r in rows:
        ws.append([r[c] for c in COLS])
    for col, wd in zip("ABCDEFGHIJK", [22, 44, 26, 34, 22, 46, 12, 36, 34, 70, 30]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(a.out)
    print(len(rows), "rows ->", a.out)


if __name__ == "__main__":
    main()
