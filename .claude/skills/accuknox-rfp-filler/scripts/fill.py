"""Stage 5, fill. Apply answers.json to the customer's workbook.

    python fill.py <workdir> <answers.json>
    python fill.py <workdir> <answers.json> --summary last     summary sheet position: first | last | none

Writes two files beside the customer's RFP:

    <name> - AccuKnox Response.xlsx   customer copy. Answers, links, evidence, branded summary sheet.
    <name> - AccuKnox Review.xlsx     internal copy. Same, plus the flag column and a Review sheet.

Never send the Review file to a customer.

Rules enforced here, per row in answers.json:
  - An existing response stays unless mode is "override" with an override_reason.
  - An existing comment is never rewritten. Links are appended, and bullets too when
    mode is "supplement". "replace_comment": true is allowed only with mode "override".
  - An existing evidence image is never replaced.
  - A NEGATIVE verdict, a draft link, a banned word or a missing link stops the run.

answers.json:
{
  "customer": "TDM Networks",
  "rows": [
    {"sheet": "SAST", "row": 9, "verdict": "MEETS", "module": "aspm",
     "bullets": ["..."], "urls": ["https://help.accuknox.com/support-matrix/sast-support-matrix/"],
     "image": "docs/..." | "gen:file.png" | "web:https://accuknox.com/...",
     "flag": "AMBER", "note": "internal reason",
     "mode": "fill" | "supplement" | "override", "override_reason": "..."}
  ]
}
"""
import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rfp_lib as L  # noqa: E402

FLAG_RANK = {"GREEN": 0, "AMBER": 1, "RED": 2}


class Stop(Exception):
    pass


def check_answer(a, sheet_prof):
    errs = []
    if a.get("verdict") == "NEGATIVE":
        errs.append("NEGATIVE verdict. Use a roadmap tier, PARTNER or PARTIAL instead.")
    if a.get("verdict") and a["verdict"] not in L.VERDICTS:
        errs.append(f"unknown verdict {a['verdict']}")
    for u in a.get("urls", []):
        if L.is_draft(u):
            errs.append(f"draft link {u}")
    text = " ".join(a.get("bullets") or [])
    hit = L.banned_hit(text)
    if hit:
        errs.append(f"banned word or em dash in bullets: {hit!r}")
    if a.get("flag") not in FLAG_RANK:
        errs.append("flag must be GREEN, AMBER or RED")
    if a.get("mode") == "override" and not a.get("override_reason"):
        errs.append("mode override needs override_reason, the user's approval in their words")
    if a.get("replace_comment") and a.get("mode") != "override":
        errs.append("replace_comment is only allowed with mode override")
    return errs


def compose(bullets, urls, existing=""):
    parts = []
    if existing:
        parts.append(existing.rstrip())
    if bullets:
        parts.append("\n".join(f"- {b.strip()}" for b in bullets))
    new_urls = [u for u in urls if u not in (existing or "")]
    if new_urls:
        parts.append("\n".join(new_urls))
    return "\n\n".join(p for p in parts if p)


def apply(wb, prof, answers, workdir, internal):
    from openpyxl.styles import Alignment, Font, PatternFill

    by_sheet = {s["sheet"]: s for s in prof["sheets"] if "rows" in s}
    log = []
    for a in answers["rows"]:
        sp = by_sheet.get(a["sheet"])
        if not sp:
            raise Stop(f"{a['sheet']}: sheet not in profile")
        rowp = next((x for x in sp["rows"] if x["row"] == a["row"]), None)
        if not rowp or rowp["kind"] != "requirement":
            raise Stop(f"{a['sheet']} r{a['row']}: not a requirement row in the profile")
        errs = check_answer(a, sp)
        if errs:
            raise Stop(f"{a['sheet']} r{a['row']}: " + "; ".join(errs))

        ws = wb[a["sheet"]]
        cols = sp["columns"]
        rc, cc, ec, nc = cols["response"], cols.get("comment"), cols.get("evidence"), cols["review_note"]
        mode = a.get("mode", "fill")
        flag = a["flag"]
        note = a.get("note", "")
        entry = {"sheet": a["sheet"], "row": a["row"], "mode": mode, "changed": []}

        # response
        cur = L.cell_text(ws.cell(a["row"], rc).value)
        entry["orig_response"] = cur or None
        if a.get("verdict") and (not cur or mode == "override"):
            resp, degraded = L.resolve_response(a["verdict"], sp["vocab"])
            if a.get("response"):
                if a["response"] not in sp["options"]:
                    raise Stop(f"{a['sheet']} r{a['row']}: response {a['response']!r} is not a dropdown option")
                resp = a["response"]
            if degraded:
                note = (note + " | " if note else "") + degraded
                if "sign-off" in degraded:
                    flag = "RED"
                elif FLAG_RANK[flag] < 1:
                    flag = "AMBER"
            if resp != cur:
                c = ws.cell(a["row"], rc)
                c.value = resp
                c.alignment = Alignment(wrap_text=True, vertical="top")
                entry["changed"].append("response")
            entry["new_response"] = resp
            if mode == "override" and cur and resp != cur:
                note = (note + " | " if note else "") + f"Changed from '{cur}'. {a['override_reason']}"
        else:
            entry["new_response"] = cur

        # comment
        if cc:
            c = ws.cell(a["row"], cc)
            existing = str(c.value or "").strip()
            if existing and not a.get("replace_comment"):
                bullets = a.get("bullets") if mode == "supplement" else None
                new = compose(bullets, a.get("urls", []), existing)
            else:
                if not a.get("bullets"):
                    raise Stop(f"{a['sheet']} r{a['row']}: empty comment cell needs bullets")
                new = compose(a["bullets"], a.get("urls", []))
            if "http" not in new and flag != "RED":
                raise Stop(f"{a['sheet']} r{a['row']}: comment has no link. Add a help.accuknox.com URL or flag RED.")
            if new != existing:
                c.value = new
                c.alignment = Alignment(wrap_text=True, vertical="top")
                entry["changed"].append("comment")

        # evidence image
        if a.get("image") and ec:
            if a["row"] in L.anchored_rows(ws):
                entry["image"] = "kept existing"
            else:
                src = L.resolve_image(a["image"], workdir)
                L.add_image(ws, a["row"], ec, src)
                entry["changed"].append("image")
                entry["image"] = a["image"]

        # review note, internal copy only
        if internal:
            c = ws.cell(a["row"], nc)
            c.value = f"{flag}. {note}".strip()
            c.fill = PatternFill("solid", fgColor=L.FLAG_FILL[flag])
            c.font = Font(color=L.FLAG_FONT[flag], size=9, bold=flag == "RED")
            c.alignment = Alignment(wrap_text=True, vertical="top")

        entry.update(flag=flag, note=note, verdict=a.get("verdict"), module=a.get("module"),
                     urls=a.get("urls", []), mandatory=rowp["mandatory"], id=rowp["id"],
                     requirement=rowp["requirement"])
        log.append(entry)

    if internal:
        for s in by_sheet.values():
            ws = wb[s["sheet"]]
            nc = s["columns"]["review_note"]
            c = ws.cell(s["header_row"], nc)
            c.value = "AccuKnox review (internal, remove before sending)"
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=L.NAVY)
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[ws.cell(1, nc).column_letter].width = 52
    return log


def summary_sheet(wb, prof, answers, log, position):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ws = wb.create_sheet("AccuKnox Summary", 0 if position == "first" else None)
    navy = PatternFill("solid", fgColor=L.NAVY)
    band = PatternFill("solid", fgColor=L.ACCENT)
    thin = Side(style="thin", color="DDE0EE")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (3, 34, 22, 22, 60)):
        ws.column_dimensions[col].width = w
    for r in range(1, 5):
        for c in range(1, 6):
            ws.cell(r, c).fill = navy
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 22
    logo = XLImage(L.LOGO)
    logo.width, logo.height = 200, 47
    ws.add_image(logo, "B2")
    customer = answers.get("customer") or "[customer name]"
    ws["E2"] = "RFP response"
    ws["E2"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["E3"] = f"Prepared for {customer}, {datetime.date.today():%d %B %Y}"
    ws["E3"].font = Font(size=11, color="D0D4F5")
    for c in range(1, 6):
        ws.cell(5, c).fill = band
    ws.row_dimensions[5].height = 4

    r = 7

    def heading(text):
        nonlocal r
        ws.cell(r, 2, text).font = Font(size=13, bold=True, color=L.NAVY)
        r += 1

    def table(headers, rows):
        nonlocal r
        for i, h in enumerate(headers):
            c = ws.cell(r, 2 + i, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = navy
            c.alignment = Alignment(vertical="center")
        r += 1
        for row in rows:
            for i, v in enumerate(row):
                c = ws.cell(r, 2 + i, v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(bottom=thin)
                if isinstance(v, str) and v.startswith("https://"):
                    c.hyperlink = v
                    c.font = Font(color=L.ACCENT, underline="single")
            r += 1
        r += 1

    heading("Response overview")
    rows = []
    for s in prof["sheets"]:
        if "rows" not in s:
            continue
        ws_c = wb[s["sheet"]]
        counts = {}
        for x in s["rows"]:
            if x["kind"] != "requirement":
                continue
            v = L.cell_text(ws_c.cell(x["row"], s["columns"]["response"]).value) or "Unanswered"
            counts[v] = counts.get(v, 0) + 1
        rows.append([s["sheet"], sum(counts.values()), "",
                     ", ".join(f"{k}: {n}" for k, n in sorted(counts.items(), key=lambda kv: -kv[1]))])
    table(["Section", "Requirements", "", "Responses"], rows)

    heading("AccuKnox modules referenced in this response")
    used = {}
    for e in log:
        if e.get("module"):
            used[e["module"]] = used.get(e["module"], 0) + 1
    mrows = []
    for mid, n in sorted(used.items(), key=lambda kv: -kv[1]):
        m = L.module_by_id(mid)
        mrows.append([m.get("short", m["name"]), n, "", m["web"][0] if m["web"] else L.doc_url(m["pages"][0])])
    table(["Module", "Rows", "", "Learn more"], mrows or [["", "", "", ""]])

    heading("How to read this response")
    for line in [
        "Every comment ends with the AccuKnox help page that documents the capability.",
        "Images in the evidence column are taken from the AccuKnox product documentation.",
        "Images titled Evidence diagram summarise a documentation page, and name that page in their footer.",
        "A roadmap answer names a capability that is planned, with the time window from the answer scale.",
    ]:
        ws.cell(r, 2, line).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        r += 1
    r += 1
    heading("Resources")
    table(["Resource", "", "", "Link"], [
        ["AccuKnox help center", "", "", "https://help.accuknox.com/"],
        ["Support matrix", "", "", "https://help.accuknox.com/support-matrix/"],
        ["Platform status", "", "", "https://status.accuknox.com/"],
        ["AccuKnox website", "", "", "https://www.accuknox.com/"],
    ])
    for c in range(1, 6):
        ws.cell(r, c).fill = band
    return ws


def review_sheet(wb, log, prof):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("AccuKnox Review", 1)
    ws.sheet_properties.tabColor = L.ALERT
    ws["A1"] = "INTERNAL. Do not send. Every filled row with its flag, verdict and evidence."
    ws["A1"].font = Font(bold=True, size=12, color=L.ALERT)
    headers = ["Flag", "Sheet", "Row", "ID", "Mandatory", "Requirement", "Response", "Verdict",
               "Module", "Changed", "Evidence image", "Links", "Review note"]
    widths = [8, 14, 6, 8, 10, 60, 22, 14, 12, 18, 40, 50, 70]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(3, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=L.NAVY)
        ws.column_dimensions[c.column_letter].width = w
    rows = sorted(log, key=lambda e: (-FLAG_RANK[e["flag"]], not e["mandatory"], e["sheet"], e["row"]))
    for n, e in enumerate(rows, 4):
        vals = [e["flag"], e["sheet"], e["row"], e["id"], "Yes" if e["mandatory"] else "",
                e["requirement"], e["new_response"], e.get("verdict"), e.get("module"),
                ", ".join(e["changed"]) or "none", e.get("image", ""), "\n".join(e["urls"]), e["note"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(n, i, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(n, 1).fill = PatternFill("solid", fgColor=L.FLAG_FILL[e["flag"]])
        ws.cell(n, 1).font = Font(bold=True, color=L.FLAG_FONT[e["flag"]])
    ws.freeze_panes = "A4"
    conflicts = L.modules().get("_conflicts", [])
    r = len(rows) + 6
    ws.cell(r, 1, "Known documentation conflicts. Resolve with product before citing.").font = Font(bold=True, color=L.ALERT)
    for i, cf in enumerate(conflicts, r + 1):
        ws.cell(i, 2, cf["topic"])
        ws.cell(i, 6, f"Older: {cf['older']}. Newer: {cf['newer']}. Rule: {cf['rule']}")
        ws.cell(i, 6).alignment = Alignment(wrap_text=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workdir")
    ap.add_argument("answers")
    ap.add_argument("--summary", choices=["first", "last", "none"], default="first")
    ap.add_argument("--out-dir")
    a = ap.parse_args()

    prof = L.read_json(os.path.join(a.workdir, "profile.json"))
    answers = L.read_json(a.answers)
    if L.sha256(prof["rfp"]) != prof["sha256"]:
        sys.exit("The customer file changed since intake. Run intake.py again.")
    out_dir = a.out_dir or os.path.dirname(prof["rfp"])
    stem = os.path.splitext(os.path.basename(prof["rfp"]))[0]
    outputs = {}
    try:
        for internal, suffix in ((True, "Review"), (False, "Response")):
            wb = L.load_workbook(prof["rfp"])
            base = sum(L.image_count(wb).values())
            log = apply(wb, prof, answers, a.workdir, internal)
            added = sum("image" in e["changed"] for e in log)
            if a.summary != "none":
                summary_sheet(wb, prof, answers, log, a.summary)
            if internal:
                review_sheet(wb, log, prof)
            path = os.path.join(out_dir, f"{stem} - AccuKnox {suffix}.xlsx")
            if L.inside_repo(path):
                sys.exit("Refusing to write a customer file inside the public help-docs repo.")
            wb.save(path)
            got = sum(L.image_count(L.load_workbook(path)).values())
            expect = base + added + (1 if a.summary != "none" else 0)
            if got < expect:
                sys.exit(f"FATAL: image loss in {path}: expected {expect}, found {got}")
            outputs[suffix] = path
    except Stop as e:
        sys.exit(f"STOPPED: {e}")

    L.write_json(os.path.join(a.workdir, "fill-manifest.json"),
                 {"outputs": outputs, "answers": os.path.abspath(a.answers), "rows": log,
                  "summary": a.summary, "images_added": added})
    flags = {k: sum(e["flag"] == k for e in log) for k in FLAG_RANK}
    print(f"filled {len(log)} rows | changed {sum(bool(e['changed']) for e in log)} | images added {added} | flags {flags}")
    for k, p in outputs.items():
        print(f"{k:9} {p}")
    print("Next: python validate.py", a.workdir)


if __name__ == "__main__":
    main()
