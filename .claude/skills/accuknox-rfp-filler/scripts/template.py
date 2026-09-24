"""Build an AccuKnox-branded response workbook.

    python template.py --blank <out.xlsx>                    empty template
    python template.py <requirements.json> <out.xlsx>        template filled with extracted requirements

Use it when the RFP arrives as Word, PDF or email text, or when the customer
sends no answer sheet. The output then runs through intake, ground, fill and
validate like any customer workbook.

requirements.json:
{
  "customer": "Acme Bank",
  "title": "Cloud security platform RFP",
  "sections": [
    {"name": "Runtime protection",
     "rows": [{"id": "RT-1", "requirement": "...", "priority": "Mandatory", "source": "RFP p.12"}]}
  ]
}
"""
import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rfp_lib as L  # noqa: E402

OPTIONS = ["1 - Meets", "2 - Exceeds", "3 - Roadmap - 30 days", "4 - Roadmap - 60 days",
           "5 - Roadmap - 90 days", "6 - Roadmap - 90+ days", "7 - Achieved via 3rd party integration",
           "8 - Partially Meets"]
HEADERS = ["ID", "Section", "Requirement", "Priority", "AccuKnox Response", "Comments", "Evidence", "Source in RFP"]
WIDTHS = [10, 22, 70, 13, 26, 70, 45, 16]


def build(req, out):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    navy = PatternFill("solid", fgColor=L.NAVY)
    band = PatternFill("solid", fgColor=L.ACCENT)
    section_fill = PatternFill("solid", fgColor="E8E9F7")
    thin = Side(style="thin", color="DDE0EE")

    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for r in (1, 2):
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c).fill = navy
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 20
    logo = XLImage(L.LOGO)
    logo.width, logo.height = 170, 40
    ws.add_image(logo, "A1")
    ws["C1"] = req.get("title") or "RFP response"
    ws["C1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["C1"].alignment = Alignment(vertical="center")
    ws["C2"] = f"Prepared for {req.get('customer') or '[customer name]'}, {datetime.date.today():%d %B %Y}"
    ws["C2"].font = Font(size=10, color="D0D4F5")
    for c in range(1, len(HEADERS) + 1):
        ws.cell(3, c).fill = band
    ws.row_dimensions[3].height = 4

    hr = 5
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(hr, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(hr + 1, 1)

    dv = DataValidation(type="list", formula1='"' + ",".join(OPTIONS) + '"', allow_blank=True)
    dv.error = "Pick a value from the list"
    ws.add_data_validation(dv)

    r = hr + 1
    for sec in req.get("sections", []):
        ws.cell(r, 3, sec["name"]).font = Font(bold=True, color=L.NAVY)
        for c in range(1, len(HEADERS) + 1):
            ws.cell(r, c).fill = section_fill
        r += 1
        for row in sec.get("rows", []):
            vals = [row.get("id", ""), sec["name"], row["requirement"], row.get("priority", ""),
                    None, None, None, row.get("source", "")]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(bottom=thin)
            dv.add(ws.cell(r, 5))
            r += 1
    if r == hr + 1:
        for _ in range(50):
            dv.add(ws.cell(r, 5))
            r += 1
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("requirements", nargs="?")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--blank", metavar="OUT")
    a = ap.parse_args()
    if a.blank:
        print("wrote", build({"title": "RFP response", "sections": []}, a.blank))
        return
    if not (a.requirements and a.out):
        sys.exit("usage: template.py <requirements.json> <out.xlsx>   or   template.py --blank <out.xlsx>")
    if L.inside_repo(a.out):
        sys.exit("Write customer templates outside the public help-docs repo.")
    req = L.read_json(a.requirements)
    n = sum(len(s.get("rows", [])) for s in req.get("sections", []))
    print(f"wrote {build(req, a.out)} with {n} requirements")
    print("Next: python intake.py", a.out)


if __name__ == "__main__":
    main()
