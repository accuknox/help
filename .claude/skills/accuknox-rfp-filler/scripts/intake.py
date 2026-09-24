"""Stage 1, pre-validation. Read any RFP workbook and describe it.

    python intake.py <rfp.xlsx>                 profile, fit check, writes profile.json
    python intake.py <rfp.docx> --extract       dump a Word RFP to text for requirement extraction
    python intake.py <rfp.xlsx> --sheet SAST    limit to named sheets (repeatable)

Detects, per sheet: the header row (including headers split over two rows), the
requirement, response, comment and evidence columns, the answer dropdown and its
vocabulary, section rows, text spillover rows, existing answers and images. Then
routes every requirement to the 12 modules and scores AccuKnox fit.
"""
import argparse
import os
import re
import sys
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rfp_lib as L  # noqa: E402

ROLE_PATTERNS = {
    "id": r"^(s\.?\s*no|sl\.?\s*no|sr\.?\s*no|#|id|ref|no\.?|item|req\.?\s*id|number)\b",
    "category": r"category|main group|group|section|domain|area|module|capability area",
    "requirement": r"requirement|specification|description|question|criteria|feature|capability|functional",
    "priority": r"mandatory|priority|desirable|must|weight|importance|criticality",
    "response": r"compliance|comply|support|status|answer|response|meets|yes\s*/\s*no",
    "comment": r"comment|detail|explanation|remark|justification|additional|notes?\b|specify|how|approach",
    "evidence": r"evidence|screenshot|proof|attachment|reference|artefact|artifact|document",
}


def _dv_options(ws, wb):
    """[(set_of_(row,col), [options])] for list-type validations."""
    out = []
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = str(dv.formula1).strip()
        opts = []
        if f.startswith('"'):
            opts = [o.strip() for o in f.strip('"').split(",") if o.strip()]
        else:
            m = re.match(r"=?'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", f)
            if m and m.group(1) in wb.sheetnames:
                src = wb[m.group(1)]
                for row in src[f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}"]:
                    for c in row:
                        if c.value is not None:
                            opts.append(str(c.value).strip())
        if not opts:
            continue
        cells = set()
        for rng in str(dv.sqref).split():
            from openpyxl.utils.cell import range_boundaries
            c1, r1, c2, r2 = range_boundaries(rng)
            for r in range(r1, min(r2, ws.max_row) + 1):
                for c in range(c1, c2 + 1):
                    cells.add((r, c))
        out.append((cells, opts))
    return out


def _header(ws, scan=30):
    """Find the header row. Merges a row with the next one when headers are split."""
    best = (0, None, {})
    for r in range(1, min(ws.max_row, scan) + 1):
        for span in (1, 2):
            labels = {}
            for rr in range(r, r + span):
                for c in range(1, min(ws.max_column, 40) + 1):
                    t = L.cell_text(ws.cell(rr, c).value)
                    if t and len(t) < 80:
                        labels[c] = (labels.get(c, "") + " " + t).strip()
            hits = {}
            for c, t in labels.items():
                for role, pat in ROLE_PATTERNS.items():
                    if re.search(pat, t, re.I):
                        hits.setdefault(c, []).append(role)
            score = len({role for roles in hits.values() for role in roles}) - (span - 1) * 0.5
            if score > best[0]:
                best = (score, (r, span), labels)
    if not best[1]:
        return None, 0, {}
    (r, span), labels = best[1], best[2]
    return r, span, labels


def _assign_columns(ws, labels, dv_cols, header_end):
    roles = {}

    def pick(role, exclude=()):
        cands = [c for c, t in labels.items() if c not in exclude and re.search(ROLE_PATTERNS[role], t, re.I)]
        return cands[0] if cands else None

    # Response: a dropdown column wins over any header text.
    if dv_cols:
        roles["response"] = dv_cols.most_common(1)[0][0]
    else:
        cands = [c for c, t in labels.items()
                 if re.search(ROLE_PATTERNS["response"], t, re.I) and not re.search(ROLE_PATTERNS["comment"], t, re.I)]
        roles["response"] = cands[0] if cands else pick("response")
    used = {roles["response"]}
    for role in ("comment", "evidence", "requirement", "id", "category", "priority"):
        c = pick(role, exclude=used)
        if role == "comment" and c is None:
            c = pick("response", exclude=used)  # "AccuKnox Response" next to a "Support" dropdown
        if c is not None:
            roles[role] = c
            used.add(c)

    # Requirement fallback: the longest text column below the header.
    if "requirement" not in roles:
        lengths = Counter()
        for r in range(header_end + 1, min(ws.max_row, header_end + 60) + 1):
            for c in range(1, min(ws.max_column, 30) + 1):
                if c not in used:
                    lengths[c] += len(L.cell_text(ws.cell(r, c).value))
        if lengths:
            roles["requirement"] = lengths.most_common(1)[0][0]
            used.add(roles["requirement"])

    if roles.get("response") and "comment" not in roles:
        nxt = roles["response"] + 1
        roles["comment"] = nxt if nxt not in used else max(used) + 1
        used.add(roles["comment"])

    # Evidence: the column most often holding an anchored image, else the next free column.
    if "evidence" not in roles:
        img_cols = Counter(im.anchor._from.col + 1 for im in ws._images)
        free = [c for c, _ in img_cols.most_common() if c not in used]
        roles["evidence"] = free[0] if free else max(used) + 1
        used.add(roles["evidence"])

    last_used = max([c for c in labels] + list(used))
    roles["review_note"] = last_used + 1
    return roles


def profile_sheet(ws, wb):
    dvs = _dv_options(ws, wb)
    head_row, span, labels = _header(ws)
    if head_row is None:
        return {"sheet": ws.title, "skipped": "no header row found"}
    header_end = head_row + span - 1

    dv_cols = Counter()
    dv_cells = set()
    options = []
    for cells, opts in dvs:
        body = {(r, c) for r, c in cells if r > header_end}
        if body:
            dv_cols.update(c for _, c in body)
            dv_cells |= body
            options = options or opts
    cols = _assign_columns(ws, labels, dv_cols, header_end)
    rc, cc, ec = cols["requirement"], cols["response"], cols.get("comment")

    vocab = L.build_vocab(options)
    img_rows = L.anchored_rows(ws)
    # With a dropdown, the answer region ends just past the last dropdown row.
    # Text below it is usually a second table with its own layout.
    region_end = max(r for r, _ in dv_cells) + 3 if dv_cells else ws.max_row
    outside = []
    if region_end < ws.max_row:
        for r, vals in enumerate(ws.iter_rows(min_row=region_end + 1, max_col=min(ws.max_column, 12),
                                              values_only=True), region_end + 1):
            if any(v not in (None, "") for v in vals):
                outside.append(r)
    req_label = labels.get(rc, "").lower()
    rows = []
    id_rate = 0
    for r in range(header_end + 1, region_end + 1):
        req = L.cell_text(ws.cell(r, rc).value)
        if not req or req.lower() == req_label:
            continue
        rid = L.cell_text(ws.cell(r, cols["id"]).value) if "id" in cols else ""
        pri = L.cell_text(ws.cell(r, cols["priority"]).value) if "priority" in cols else ""
        resp = L.cell_text(ws.cell(r, cc).value)
        com = str(ws.cell(r, ec).value or "").strip() if ec else ""
        other = [L.cell_text(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 30) + 1) if c != rc]
        rows.append(dict(row=r, id=rid, priority=pri, requirement=req, response=resp,
                         comment_len=len(com), comment_has_url="http" in com,
                         has_image=r in img_rows, in_dropdown=(r, cc) in dv_cells,
                         others=sum(1 for x in other if x)))
        id_rate += bool(rid)
    id_rate = id_rate / max(1, len(rows))

    section = None
    out = []
    for i, x in enumerate(rows):
        kind = "requirement"
        bare = not x["id"] and not x["priority"] and not x["response"]
        if bare and dv_cells and not x["in_dropdown"] and len(x["requirement"]) < 120:
            kind = "section"
        elif bare and x["others"] == 0 and len(x["requirement"]) < 90:
            kind = "section"
        if kind == "requirement" and id_rate > 0.6 and not x["id"] and not x["priority"] and i > 0:
            prev = out[-1] if out else None
            if prev and prev["kind"] == "requirement" and not re.search(r"[.?!:;)]$", prev["requirement"]):
                kind = "spillover"
        if kind == "section":
            section = x["requirement"]
        context = f"{ws.title} {section or ''}"
        routed = L.route(x["requirement"], context) if kind == "requirement" else []
        out.append(dict(
            row=x["row"], kind=kind, id=x["id"], section=section if kind != "section" else None,
            priority=x["priority"], mandatory=bool(re.search(r"mandatory|must|required|\bm\b|high", x["priority"] + " " + x["requirement"][:15], re.I)),
            requirement=x["requirement"], existing_response=x["response"] or None,
            existing_verdict=L.classify_option(x["response"]) if x["response"] else None,
            comment_len=x["comment_len"], comment_has_url=x["comment_has_url"], has_image=x["has_image"],
            modules=[m for m, _ in routed], route_score=routed[0][1] if routed else 0,
            direct=bool(L.route(x["requirement"])) if kind == "requirement" else False,
            fit_risks=L.fit_risks(x["requirement"]) if kind == "requirement" else []))

    reqs = [x for x in out if x["kind"] == "requirement"]
    return {
        "sheet": ws.title, "header_row": head_row, "header_rows": span,
        "headers": {str(k): v for k, v in labels.items()},
        "columns": cols, "options": options, "vocab": vocab,
        "unmapped_options": [o for o in options if not L.classify_option(o)],
        "has_negative_option": "NEGATIVE" in vocab,
        "has_roadmap_option": any(v.startswith("ROADMAP") for v in vocab),
        "images": len(ws._images),
        "answer_region": [header_end + 1, region_end],
        "outside_rows": {"count": len(outside), "first": outside[0] if outside else None,
                         "last": outside[-1] if outside else None},
        "counts": {
            "requirements": len(reqs),
            "sections": sum(x["kind"] == "section" for x in out),
            "spillover": sum(x["kind"] == "spillover" for x in out),
            "answered": sum(bool(x["existing_response"]) for x in reqs),
            "blank": sum(not x["existing_response"] for x in reqs),
            "negative_existing": sum(x["existing_verdict"] == "NEGATIVE" for x in reqs),
            "no_url": sum(x["comment_len"] > 0 and not x["comment_has_url"] for x in reqs),
            "no_image": sum(not x["has_image"] for x in reqs),
            "mandatory": sum(x["mandatory"] for x in reqs),
        },
        "rows": out,
    }


def fit(sheets):
    reqs = [x for s in sheets if "rows" in s for x in s["rows"] if x["kind"] == "requirement"]
    if not reqs:
        return {"verdict": "UNKNOWN", "reason": "no requirement rows found"}
    routed = [x for x in reqs if x["direct"]]
    risky = [x for x in reqs if x["fit_risks"]]
    mand = [x for x in reqs if x["mandatory"]]
    mand_risky = [x for x in mand if x["fit_risks"] or not x["direct"]]
    cov = len(routed) / len(reqs)
    by_mod = Counter(m for x in reqs for m in x["modules"][:1])
    if cov >= 0.75 and len(mand_risky) <= max(1, 0.05 * len(mand)):
        verdict = "GO"
    elif cov >= 0.5 and len(mand_risky) <= 0.25 * max(1, len(mand)):
        verdict = "CONDITIONAL"
    else:
        verdict = "NO-GO"
    return {
        "verdict": verdict,
        "coverage": round(cov, 3),
        "requirements": len(reqs),
        "routed": len(routed),
        "unrouted_rows": [(x.get("_sheet"), x["row"], x["requirement"][:90]) for x in reqs if not x["direct"]][:40],
        "fit_risk_rows": [(x.get("_sheet"), x["row"], x["fit_risks"], x["mandatory"]) for x in risky][:40],
        "mandatory": len(mand),
        "mandatory_at_risk": len(mand_risky),
        "primary_modules": by_mod.most_common(),
    }


def extract_docx(path, out_txt):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    lines = []
    body = root.find("w:body", ns)
    for el in body:
        tag = el.tag.split("}")[1]
        if tag == "p":
            t = "".join(x.text or "" for x in el.iter(f"{{{ns['w']}}}t")).strip()
            if t:
                lines.append(t)
        elif tag == "tbl":
            for tr in el.iter(f"{{{ns['w']}}}tr"):
                cells = ["".join(x.text or "" for x in tc.iter(f"{{{ns['w']}}}t")).strip()
                         for tc in tr.iter(f"{{{ns['w']}}}tc")]
                lines.append(" | ".join(cells))
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return len(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("rfp")
    ap.add_argument("--sheet", action="append")
    ap.add_argument("--extract", action="store_true")
    ap.add_argument("--allow-repo", action="store_true", help="only for the skill self-test")
    a = ap.parse_args()

    if not os.path.isfile(a.rfp):
        sys.exit(f"not found: {a.rfp}")
    wd = L.workdir_for(a.rfp, allow_repo=a.allow_repo)
    ext = os.path.splitext(a.rfp)[1].lower()

    if a.extract or ext == ".docx":
        if ext != ".docx":
            sys.exit("--extract handles .docx. Read a PDF with the Read tool, page by page.")
        out = os.path.join(wd, "rfp-text.txt")
        n = extract_docx(a.rfp, out)
        print(f"extracted {n} lines -> {out}")
        print("Next: turn the requirements into requirements.json and build the AccuKnox template with template.py.")
        return
    if ext not in (".xlsx", ".xlsm"):
        sys.exit(f"unsupported file type {ext}. Use .xlsx, .xlsm or .docx, or read a PDF with the Read tool.")

    wb = L.load_workbook(a.rfp)
    sheets = []
    for ws in wb.worksheets:
        if a.sheet and ws.title not in a.sheet:
            continue
        if ws.sheet_state != "visible":
            continue
        s = profile_sheet(ws, wb)
        for x in s.get("rows", []):
            x["_sheet"] = ws.title
        sheets.append(s)

    blockers = []
    for s in sheets:
        if "skipped" in s:
            continue
        if not s["options"]:
            blockers.append(f"{s['sheet']}: no answer dropdown found. Confirm the response column and vocabulary with the user.")
        elif "MEETS" not in s["vocab"] and "EXCEEDS" not in s["vocab"]:
            blockers.append(f"{s['sheet']}: no positive option recognised in {s['options']}. Map it by hand.")
        if s["unmapped_options"]:
            blockers.append(f"{s['sheet']}: unmapped dropdown options {s['unmapped_options']}.")
        if s["outside_rows"]["count"]:
            o = s["outside_rows"]
            blockers.append(f"{s['sheet']}: {o['count']} non-empty rows below the answer range (rows {o['first']}-{o['last']}). "
                            "Likely a second table with its own layout. Ask the user whether it needs answers.")
        if not s["has_roadmap_option"]:
            blockers.append(f"{s['sheet']}: no roadmap option. Roadmap rows need a human decision.")

    prof = {
        "rfp": os.path.abspath(a.rfp), "sha256": L.sha256(a.rfp), "workdir": wd,
        "images_baseline": L.image_count(wb), "sheets": sheets,
        "fit": fit(sheets), "blockers": blockers,
    }
    L.write_json(os.path.join(wd, "profile.json"), prof)

    print(f"RFP      {a.rfp}")
    print(f"workdir  {wd}")
    for s in sheets:
        if "skipped" in s:
            print(f"  [{s['sheet']}] skipped: {s['skipped']}")
            continue
        c = s["counts"]
        colmap = ", ".join(f"{k}={chr(64 + v)}" for k, v in s["columns"].items() if v and v <= 26)
        print(f"  [{s['sheet']}] header row {s['header_row']}(+{s['header_rows'] - 1}) | {colmap}")
        print(f"      requirements {c['requirements']} | blank {c['blank']} | mandatory {c['mandatory']} | "
              f"sections {c['sections']} | spillover {c['spillover']} | images {s['images']} | "
              f"no image {c['no_image']} | comment without link {c['no_url']} | existing negative {c['negative_existing']}")
        print(f"      vocabulary {s['vocab']}")
    f = prof["fit"]
    print(f"FIT      {f['verdict']} | coverage {f.get('coverage')} | mandatory at risk {f.get('mandatory_at_risk')}/{f.get('mandatory')}")
    print(f"modules  {f.get('primary_modules')}")
    for r in f.get("fit_risk_rows", [])[:15]:
        print(f"  risk   {r}")
    for b in blockers:
        print(f"BLOCKER  {b}")
    print(f"wrote    {os.path.join(wd, 'profile.json')}")


if __name__ == "__main__":
    main()
