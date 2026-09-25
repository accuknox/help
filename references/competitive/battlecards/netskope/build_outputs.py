# -*- coding: utf-8 -*-
"""Build the XLSX workbook and the markdown draft from netskope_data.py.

Run from the repo root:
    py -3.11 references/competitive/battlecards/netskope/build_outputs.py
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import netskope_data as d  # noqa: E402

REPO = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
XLSX = os.path.join(REPO, "references", "competitive", "battlecards", "accuknox-vs-netskope-ai-security.xlsx")
MD = os.path.join(REPO, "references", "drafts", "comparison", "accuknox-vs-netskope-ai-security.md")

FONT = "Space Grotesk"
NAVY, PRIMARY, INK, MUTE = "11206D", "0046FF", "1B223B", "5A637D"
LAV, GREY_BG, GREY_BD = "ECECFB", "EEF0F6", "C4CCDE"
GREEN_LT, GREEN_DK = "E7F6EE", "0B7A42"
RED_LT, RED = "FBECEE", "C80019"
WHITE = "FFFFFF"

thin = Side(style="thin", color=GREY_BD)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def f(bold=False, color=INK, size=10, italic=False, underline=None):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


STATUS_FILL = {
    "GA": (GREEN_LT, GREEN_DK),
    "Beta": (LAV, "4D4DD9"),
    "Preview": (LAV, "4D4DD9"),
    "Roadmap": (GREY_BG, MUTE),
    "Not documented": (GREY_BG, MUTE),
}
EDGE_FILL = {
    d.AK: (GREEN_LT, GREEN_DK),
    d.NS: (GREY_BG, NAVY),
    d.PAR: (LAV, "4D4DD9"),
    d.NA: (GREY_BG, MUTE),
}


def status_style(cell, value):
    key = "GA" if value.startswith("GA") else value
    bg, fg = STATUS_FILL.get(key, (WHITE, INK))
    if "Beta" in value and key == "GA":
        bg, fg = LAV, "4D4DD9"
    cell.fill = fill(bg)
    cell.font = f(bold=True, color=fg)
    cell.alignment = CENTER


def header_row(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = f(bold=True, color=WHITE)
        c.fill = fill(NAVY)
        c.alignment = CENTER
        c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def src_text(keys):
    return "\n".join(d.SOURCES[k][2] for k in keys)


# ============================================================== XLSX
def build_xlsx():
    wb = Workbook()

    # ---------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 70

    ws["A1"] = f"{d.AK_NAME} vs {d.NS_NAME}"
    ws["A1"].font = f(bold=True, color=NAVY, size=18)
    ws["A2"] = (f"Technical capability comparison. Sources read on {d.READ_ON}. "
                "Every cell cites help.accuknox.com, accuknox.com, docs.netskope.com or netskope.com.")
    ws["A2"].font = f(color=MUTE, size=10, italic=True)
    ws.merge_cells("A1:E1")
    ws.merge_cells("A2:E2")

    sc = d.score()
    r = 4
    header_row(ws, r, ["Scorecard", "AccuKnox edge", "Netskope edge", "Parity", "What decides the row"])
    r += 1
    for cid, cname in d.CATEGORIES:
        rows = [x for x in d.ROWS if x["cat"] == cid]
        vals = [cname,
                sum(1 for x in rows if x["edge"] == d.AK),
                sum(1 for x in rows if x["edge"] == d.NS),
                sum(1 for x in rows if x["edge"] == d.PAR),
                " ".join(x["capability"] + "." for x in rows)]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = f(bold=(i == 1))
            c.alignment = CENTER if i in (2, 3, 4) else WRAP
            c.border = BORDER
        r += 1
    tot = ["Total, 16 capabilities", sc[d.AK], sc[d.NS], sc[d.PAR],
           f"{sc[d.NA]} row not scored: licensing, until AccuKnox confirms its licensing unit."]
    for i, v in enumerate(tot, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = f(bold=True, color=WHITE if i < 5 else NAVY)
        c.fill = fill(PRIMARY if i < 5 else LAV)
        c.alignment = CENTER if i in (2, 3, 4) else WRAP
        c.border = BORDER
    r += 2

    header_row(ws, r, ["Where Netskope leads", "", "", "", "Fact and source"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    for head, text, key in d.NETSKOPE_LEADS:
        ws.cell(row=r, column=1, value=head).font = f(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=5, value=f"{text}\n{d.SOURCES[key][2]}")
        c.alignment = WRAP
        c.font = f()
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BORDER
        r += 1
    r += 1

    header_row(ws, r, ["Why AccuKnox", "", "", "", "Fact and source"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    for head, text, keys in d.BETTER_FASTER_CHEAPER:
        ws.cell(row=r, column=1, value=head).font = f(bold=True, color=PRIMARY)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=5, value=text + "\n" + src_text(keys))
        c.alignment = WRAP
        c.font = f()
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BORDER
        r += 1
    r += 1

    header_row(ws, r, ["Status legend", "", "", "", "Meaning"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    r += 1
    legend = [
        ("GA", "Generally available, documented by the vendor."),
        ("Beta", "The vendor marks it Beta. AccuKnox AI GRC and AI Identity Security are Beta."),
        ("Preview", "Netskope marks it a preview feature."),
        ("Roadmap", "Announced, not in the product today."),
        ("Not documented", "The vendor's AI Security docs and product pages do not describe it."),
        ("GA, separate license", "Available, but needs a license beyond the base AI product."),
    ]
    for k, v in legend:
        c = ws.cell(row=r, column=1, value=k)
        status_style(c, k)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c2 = ws.cell(row=r, column=5, value=v)
        c2.font = f()
        c2.alignment = WRAP
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = BORDER
        r += 1

    # ------------------------------------------------------- Comparison
    ws = wb.create_sheet("Comparison")
    ws.sheet_view.showGridLines = False
    cols = ["ID", "Capability", "AccuKnox: what is supported and how", "AccuKnox status",
            "AccuKnox sources", "Netskope: what is supported and how", "Netskope status",
            "Netskope sources", "Edge", "Why"]
    header_row(ws, 1, cols, [6, 24, 62, 14, 42, 62, 14, 42, 12, 40])
    ws.freeze_panes = "C2"
    r = 2
    for cid, cname in d.CATEGORIES:
        c = ws.cell(row=r, column=1, value=f"{cid}.  {cname}")
        c.font = f(bold=True, color=PRIMARY, size=11)
        for col in range(1, len(cols) + 1):
            ws.cell(row=r, column=col).fill = fill(LAV)
            ws.cell(row=r, column=col).border = BORDER
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
        ws.row_dimensions[r].height = 22
        r += 1
        for x in [x for x in d.ROWS if x["cat"] == cid]:
            vals = [x["id"], x["capability"], x["ak"], x["ak_status"], src_text(x["ak_src"]),
                    x["ns"], x["ns_status"], src_text(x["ns_src"]), x["edge"], x["why"]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = f(bold=(i in (1, 2)), size=10, color=(NAVY if i in (1, 2) else INK))
                c.alignment = WRAP
                c.border = BORDER
            for i in (5, 8):
                c = ws.cell(row=r, column=i)
                c.font = f(size=9, color=PRIMARY, underline="single")
                c.hyperlink = c.value.split("\n")[0]
            status_style(ws.cell(row=r, column=4), x["ak_status"])
            status_style(ws.cell(row=r, column=7), x["ns_status"])
            bg, fg = EDGE_FILL[x["edge"]]
            ec = ws.cell(row=r, column=9)
            ec.fill = fill(bg)
            ec.font = f(bold=True, color=fg)
            ec.alignment = CENTER
            lines = max(len(x["ak"]), len(x["ns"])) // 58 + 2
            ws.row_dimensions[r].height = max(60, 14 * lines)
            r += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{r - 1}"

    # ---------------------------------------------------- Beta and stage
    ws = wb.create_sheet("Beta and Preview")
    ws.sheet_view.showGridLines = False
    header_row(ws, 1, ["Vendor", "Capability", "Stage", "What the source says", "Source"],
               [14, 34, 14, 70, 60])
    stage = [
        ("AccuKnox", "AI Identity Security", "Beta", "SPIFFE identity per agent, OpenFGA permissions, caller chain check.", "ak-web-ai"),
        ("AccuKnox", "AI GRC", "Beta", "EU AI Act risk tiers, ISO 42001, 12+ frameworks, audit reports.", "ak-web-grc"),
        ("AccuKnox", "MCP protocol policies", "Roadmap", "MCP discovery and process sandbox ship today. Protocol-level MCP policy is roadmap.", "ak-v35"),
        ("Netskope", "Agent Action Control", "Beta", "Formerly Agent Guardrails. Support enables it per tenant.", "ns-aac"),
        ("Netskope", "AI Guardrails custom topics", "Beta", "Up to 30 custom topics. Support enables it per tenant.", "ns-topics"),
        ("Netskope", "AI Platforms discovery and posture", "Preview", "13 AI platforms on AWS, Azure and GCP.", "ns-platforms"),
        ("Netskope", "Anthropic Claude support for AICC", "Preview", "Claude Enterprise and Claude Platform posture through the Compliance and Admin APIs.", "ns-claude"),
        ("Netskope", "Agentic Broker", "GA, separate license", "Account team enables the license. DLP on MCP needs the DLP add-on.", "ns-broker"),
    ]
    for i, (v, cap, st, txt, key) in enumerate(stage, 2):
        vals = [v, cap, st, txt, d.SOURCES[key][2]]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = f(bold=(j == 2))
            c.alignment = WRAP
            c.border = BORDER
        status_style(ws.cell(row=i, column=3), st)
        lc = ws.cell(row=i, column=5)
        lc.hyperlink = lc.value
        lc.font = f(size=9, color=PRIMARY, underline="single")

    # ------------------------------------------------ Live page fixes
    ws = wb.create_sheet("Live Page Fixes")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Claims on accuknox.com/comparisons/accuknox-vs-netskope that the sources contradict"
    ws["A1"].font = f(bold=True, color=NAVY, size=13)
    ws.merge_cells("A1:D1")
    header_row(ws, 2, ["#", "Live page says", "The source says", "Source"], [5, 60, 70, 60])
    for i, (claim, fact, key) in enumerate(d.LIVE_PAGE_FIXES, 3):
        vals = [i - 2, claim, fact, d.SOURCES[key][2]]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = f()
            c.alignment = WRAP
            c.border = BORDER
        lc = ws.cell(row=i, column=4)
        lc.hyperlink = lc.value
        lc.font = f(size=9, color=PRIMARY, underline="single")
        ws.row_dimensions[i].height = 44

    # ---------------------------------------------------------- Sources
    ws = wb.create_sheet("Sources")
    ws.sheet_view.showGridLines = False
    header_row(ws, 1, ["Key", "Vendor", "Page", "URL", "Used in rows"], [16, 12, 44, 90, 22])
    used = {}
    for x in d.ROWS:
        for k in x["ak_src"] + x["ns_src"]:
            used.setdefault(k, []).append(x["id"])
    for i, (k, (v, title, url)) in enumerate(d.SOURCES.items(), 2):
        vals = [k, v, title, url, ", ".join(used.get(k, [])) or "Summary, fixes"]
        for j, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = f()
            c.border = BORDER
            c.alignment = WRAP
        lc = ws.cell(row=i, column=4)
        lc.hyperlink = url
        lc.font = f(size=9, color=PRIMARY, underline="single")

    wb.save(XLSX)
    return XLSX


# ============================================================== MD
def link(key):
    v, title, url = d.SOURCES[key]
    return f"[{title}]({url})"


def build_md():
    sc = d.score()
    out = ["---",
           'title: "AccuKnox (vs) Netskope"',
           'subtitle: "AI Security Platform Comparison"',
           'slug: "accuknox-vs-netskope"',
           'url: "https://accuknox.com/comparisons/accuknox-vs-netskope"',
           'archetype: "head-to-head"',
           'category: "ai-security"',
           'competitors: ["Netskope Skylight AI Security"]',
           f"parameter_count: {len(d.ROWS)}",
           'excerpt: "AccuKnox AI Security and Netskope Skylight AI Security compared on 16 sourced '
           'capabilities across posture, models, guardrails, agents, response and deployment."',
           "---", ""]
    out.append("# AccuKnox (vs) Netskope\n")
    out.append("## AI Security Platform Comparison\n")
    out.append(
        f"AccuKnox leads {sc[d.AK]} of 16 sourced rows and Netskope leads {sc[d.NS]}. "
        "The rows cover AI posture, model security, red teaming, guardrails, agents and MCP, response, governance and deployment. "
        "Netskope leads where the job is workforce traffic, because its proxy carries every user request. "
        "AccuKnox leads where the job is the model, the agent host and the cloud AI service. "
        f"Every cell below cites the vendor's own documentation, read on {d.READ_ON}.\n")
    out.append("AccuKnox marks AI GRC and AI Identity Security as Beta. "
               "Netskope marks Agent Action Control and custom guardrail topics as Beta, and AI Platforms posture as a preview feature.\n")
    for cid, cname in d.CATEGORIES:
        out.append(f"## {d.HEADLINES[cid]}\n")
        out.append("| Parameter | AccuKnox | Netskope |")
        out.append("| --- | --- | --- |")
        for x in [x for x in d.ROWS if x["cat"] == cid]:
            ak = f"**{x['ak_status']}.** {x['ak']} References: " + ", ".join(link(k) for k in x["ak_src"])
            ns = f"**{x['ns_status']}.** {x['ns']} References: " + ", ".join(link(k) for k in x["ns_src"])
            out.append(f"| **{x['capability']}** | {ak} | {ns} |")
        out.append("")
    out.append("## Why Customers Choose AccuKnox Over Netskope\n")
    for head, text, keys in d.BETTER_FASTER_CHEAPER:
        out.append(f"### {head}\n")
        out.append(text + " References: " + ", ".join(link(k) for k in keys) + "\n")
    out.append("## Where Netskope Leads Today\n")
    for head, text, key in d.NETSKOPE_LEADS:
        out.append(f"- **{head}.** {text} Source: {link(key)}.")
    out.append("")
    return "\n".join(out)


if __name__ == "__main__":
    print("xlsx:", build_xlsx())
    os.makedirs(os.path.dirname(MD), exist_ok=True)
    with open(MD, "w", encoding="utf-8") as fh:
        fh.write(build_md())
    print("md:", MD)
