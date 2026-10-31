"""Build the short Shadow AI Discovery handoff PDF and the paste-ready sheet.

Usage:
    python build.py

Outputs, next to this file: shadow-ai-comparison-updates.pdf, .html, .xlsx, .csv, .md
"""

from __future__ import annotations

import csv
import html
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from content import VENDORS  # noqa: E402

OUT = Path(__file__).parent
STEM = "shadow-ai-comparison-updates"
CLS = {"Yes": "ok", "Dash": "part", "No": "no"}


def e(s: str) -> str:
    return html.escape(s, quote=False)


SVG = {
    "Yes": '<svg viewBox="0 0 24 24" width="15" height="15"><path d="M4 12.5l5 5L20 6.5" fill="none" stroke="#16B030" stroke-width="3.2" stroke-linecap="round" stroke-linejoin="round"/></svg>',
    "Dash": '<svg viewBox="0 0 24 24" width="15" height="15"><path d="M6 12h12" stroke="#111" stroke-width="3.2" stroke-linecap="round"/></svg>',
    "No": '<svg viewBox="0 0 24 24" width="15" height="15"><path d="M6 6l12 12M18 6L6 18" stroke="#E02424" stroke-width="3.2" stroke-linecap="round"/></svg>',
}


def chip(icon: str) -> str:
    return f'<span class="ic">{SVG[icon]}</span>'


def short(url: str) -> str:
    return url.replace("https://", "")


CSS = """
@page { size: A4; margin: 12mm 13mm 13mm 13mm; }
* { box-sizing: border-box; margin: 0; padding: 0; }
:root { --navy:#11206D; --blue:#0046FF; --ink:#0F172A; --body:#334155; --muted:#64748B; --line:#E2E8F0;
  --soft:#F4F7FF; --ok:#0E9F6E; --part:#D97706; --no:#DC2626; --rev:#B45309; }
body { font-family: Inter, "Segoe UI", Arial, sans-serif; font-size: 8.6pt; line-height: 1.4; color: var(--body);
  -webkit-print-color-adjust: exact; print-color-adjust: exact; }
a { color: var(--blue); text-decoration: none; }
.top { background: var(--navy); color:#fff; border-radius: 2.5mm; padding: 3mm 4.5mm; margin-bottom: 3mm; }
.top h1 { color:#fff; font-size: 13pt; letter-spacing: -.01em; }
.top p { color:#C9D3F5; font-size: 8.4pt; margin-top: 1mm; }
.legend { margin-top: 2mm; font-size: 7.8pt; color:#C9D3F5; display:flex; gap: 4mm; flex-wrap: wrap; }
.chip { display:inline-block; font-size: 6.8pt; font-weight: 700; padding: .3mm 1.6mm; border-radius: 99px; color:#fff; }
.chip.ok { background: var(--ok); } .chip.part { background: var(--part); } .chip.no { background: var(--no); }
.chip.rev { background:#FEF3C7; color: var(--rev); border: 1px solid #F59E0B; }
.v { border: 1px solid var(--line); border-radius: 2.5mm; padding: 3mm 4mm; margin-bottom: 3mm; page-break-inside: avoid; }
.vh { display:flex; justify-content: space-between; align-items: baseline; }
.vh b { font-size: 11pt; color: var(--navy); }
.vh a { font-size: 7.4pt; }
.where { font-size: 7.6pt; color: var(--muted); margin-top: .6mm; }
.where b { color: var(--blue); }
.h2 { font-size: 11.5pt; font-weight: 800; color: var(--ink); margin-top: 2.4mm; line-height: 1.2; }
.intro { color: var(--ink); margin-top: 1.4mm; }
table { width: 100%; border-collapse: collapse; margin-top: 2.4mm; font-size: 7.8pt; }
th { text-align: left; background: var(--soft); color: var(--ink); font-size: 7pt; padding: 1.4mm 1.8mm; }
td { padding: 1.2mm 1.8mm; border-top: 1px solid var(--line); vertical-align: top; }
td.r { font-weight: 700; color: var(--ink); width: 26%; }
td.a, td.c { width: 37%; }
td .t { display:block; margin-top: .6mm; }
.ic { display:inline-block; vertical-align: middle; }
.ic svg { display:block; }
td .ref { display:block; font-size: 6.4pt; margin-top: .3mm; line-height:1.25; }
.src { font-size: 6.6pt; color: var(--muted); margin-top: 1.8mm; }
.end { page-break-inside: avoid; border: 1px solid var(--line); border-radius: 2.5mm; padding: 3.5mm 4.5mm; }
.end h2 { font-size: 11pt; color: var(--navy); margin-bottom: 1.4mm; }
.end .sh { font-weight: 700; color: var(--ink); margin-top: 2mm; font-size: 8pt; }
.end li { list-style: none; padding: .4mm 0; font-size: 7.4pt; }
"""


def vendor_html(v: dict) -> str:
    rows = ""
    for sub, (ai, at, aref), (ci, ct, rev) in v["rows"]:
        r = ' <span class="chip rev">REVIEW</span>' if rev else ""
        rows += (f"<tr><td class='r'>Shadow AI - {e(sub)}</td>"
                 f"<td class='a'>{chip(ai)}<span class='t'>{e(at)}</span><span class='ref'>Reference: <a href='{aref}'>{e(short(aref))}</a></span></td>"
                 f"<td class='c'>{chip(ci)}{r}<span class='t'>{e(ct)}</span></td></tr>")
    return f"""
<div class="v">
  <div class="vh"><b>AccuKnox vs {e(v['name'])}</b><a href="{v['url']}">{e(short(v['url']))}</a></div>
  <div class="where"><b>Where:</b> {e(v['where'])}</div>
  <table><thead><tr><th>ROW</th><th>ACCUKNOX</th><th>{e(v['name'].upper())}</th></tr></thead><tbody>{rows}</tbody></table>
</div>"""


def build_html() -> str:
    top = f"""
<div class="top"><h1>Shadow AI Discovery Blocks for Six Comparison Pages</h1>
<p>Add the 4 rows to each page's comparison table. Use the same icon in the Image Select field. Hold REVIEW cells until approved.</p></div>"""
    ak_refs = sorted({a[2] for v in VENDORS for _, a, _ in v["rows"]})
    end = ("<div class='end'><h2>Sources and References</h2><p class='sh'>AccuKnox</p><ul>"
           + "<li>AccuKnox Shadow AI Discovery one-pager, Oct 2026</li>"
           + "".join(f"<li><a href='{u}'>{e(short(u))}</a></li>" for u in ak_refs) + "</ul>"
           + "".join(f"<p class='sh'>{e(v['name'])} (read 2026-10-27)</p><ul>"
                     + "".join(f"<li><a href='https://{s}'>{e(s)}</a></li>" for s in v["sources"]) + "</ul>" for v in VENDORS)
           + "</div>")
    body = top + "".join(vendor_html(v) for v in VENDORS) + end
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>Shadow AI Comparison Updates</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap" rel="stylesheet">
<style>{CSS}</style></head><body>{body}</body></html>"""


def review_list(v: dict) -> str:
    items = [f"{sub}: {ct}" for sub, _, (ci, ct, rev) in v["rows"] if rev]
    return "\n".join(f"- {x}" for x in items) or "None"


def write_sheet() -> None:
    header = ["Vendor", "Where", "Row", "AccuKnox icon", "AccuKnox text", "Reference link", "Competitor icon", "Competitor text", "Review"]
    data = [[v["name"], v["where"], f"Shadow AI - {sub}", ai, at, aref, ci, ct, "REVIEW" if rev else ""]
            for v in VENDORS for sub, (ai, at, aref), (ci, ct, rev) in v["rows"]]
    with open(OUT / f"{STEM}.csv", "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows([header] + data)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Table rows"
    ws.append(header)
    for d in data:
        ws.append(d)
    for i, w in enumerate([12, 40, 34, 10, 55, 45, 10, 50, 10], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="11206D")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUT / f"{STEM}.xlsx")


def write_md() -> None:
    out = []
    for v in VENDORS:
        out += [f"## {v['name']}", ""]
        for sub, a, c in v["rows"]:
            out += [f"### {sub}", "", a[1], "", c[1], ""]
    (OUT / f"{STEM}.md").write_text("\n".join(out), encoding="utf-8")


def write_pdf(hp: Path) -> None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.goto(hp.as_uri(), wait_until="networkidle")
        pg.pdf(path=str(OUT / f"{STEM}.pdf"), prefer_css_page_size=True, print_background=True)
        b.close()


if __name__ == "__main__":
    hp = OUT / f"{STEM}.html"
    hp.write_text(build_html(), encoding="utf-8")
    write_sheet()
    write_md()
    write_pdf(hp)
    print("built", OUT / f"{STEM}.pdf")
